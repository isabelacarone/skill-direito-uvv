"""Gera os .docx de teste do modo Checar documento (conteudo ficticio).

Uso: python tests/checar/gerar_fixtures.py  (grava ao lado deste script)
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CORPO = (
    "Este paragrafo de teste existe apenas para exercitar o verificador de formatacao e nao "
    "tem qualquer valor juridico. Ele repete uma frase ficticia ate passar de duzentos "
    "caracteres, para que o texto ocupe varias linhas e seja tratado como corpo do documento "
    "durante a checagem de fonte, tamanho, entrelinhas e alinhamento."
)
CITACAO = (
    "Esta citacao longa e inteiramente ficticia e serve para testar a exclusao de citacoes "
    "do corpo do texto. Ela precisa ter mais de trezentos caracteres, sem aspas, em paragrafo "
    "isolado e com recuo de quatro centimetros da margem esquerda, em fonte menor e com "
    "espacamento simples, como pede a regra de citacao direta com mais de tres linhas, de "
    "modo que o verificador a reconheca como candidata a citacao longa. (Teste, 2020, p. 10)."
)


def _fonte_run(run, fonte, tamanho, negrito=None):
    run.font.name = fonte
    run.font.size = Pt(tamanho)
    if negrito is not None:
        run.font.bold = negrito
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    rfonts.set(qn("w:hAnsi"), fonte)


def _par(doc, texto, fonte, tamanho=12, alin=None, negrito=None, recuo_cm=None,
         simples=False, quebra=False):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    _fonte_run(r, fonte, tamanho, negrito)
    if alin is not None:
        p.alignment = alin
    if recuo_cm is not None:
        p.paragraph_format.left_indent = Cm(recuo_cm)
    if simples:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if quebra:
        p.paragraph_format.page_break_before = True
    return p


def gerar_base(caminho, margem_cm=None, fonte="Arial"):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = fonte
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:hAnsi"), fonte)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        if margem_cm is None:
            s.top_margin, s.left_margin = Cm(3), Cm(3)
            s.right_margin, s.bottom_margin = Cm(2), Cm(2)
        else:
            s.top_margin = s.left_margin = Cm(margem_cm)
            s.right_margin = s.bottom_margin = Cm(margem_cm)
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi"):
            rf.set(qn(a), "Arial")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "20")
        rpr.append(rf)
        rpr.append(sz)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(rpr)
        r.append(t)
        fld.append(r)
        hp._p.append(fld)

    C = WD_ALIGN_PARAGRAPH.CENTER
    L = WD_ALIGN_PARAGRAPH.LEFT
    for txt in ("UNIVERSIDADE DE TESTE", "MARIA DE TESTE", "TÍTULO DE TESTE",
                "VILA VELHA-ES", "2026"):
        _par(doc, txt, fonte, alin=C)

    _par(doc, "1 INTRODUÇÃO", fonte, alin=L, negrito=True, quebra=True)
    for _ in range(3):
        _par(doc, CORPO, fonte)
    _par(doc, "1.1 SEÇÃO SECUNDÁRIA DE TESTE", fonte, alin=L, negrito=False)
    _par(doc, "1.1.1 Seção terciária de teste", fonte, alin=L, negrito=True)
    _par(doc, CORPO, fonte)
    _par(doc, CITACAO, "Arial", tamanho=10, alin=WD_ALIGN_PARAGRAPH.JUSTIFY,
         recuo_cm=4, simples=True)
    _par(doc, "2 SEGUNDA SEÇÃO DE TESTE", fonte, alin=L, negrito=True, quebra=True)
    _par(doc, CORPO, fonte)
    _par(doc, "REFERÊNCIAS", fonte, alin=C, negrito=True, quebra=True)
    _par(doc, "TESTE, Maria de. Obra de teste. Vila Velha: Editora de Teste, 2020.",
         fonte, alin=L, simples=True)
    doc.save(str(caminho))


def gerar_caso01(caminho):
    gerar_base(caminho, margem_cm=2.5, fonte="Times New Roman")


def gerar_caso03(caminho):
    gerar_base(caminho)


def main():
    aqui = Path(__file__).resolve().parent
    gerar_caso01(aqui / "caso-01-margem-fonte.docx")
    gerar_caso03(aqui / "caso-03-conforme.docx")
    print("fixtures gravadas em", aqui)


if __name__ == "__main__":
    sys.exit(main())

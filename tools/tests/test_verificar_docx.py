import importlib.util
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.shared import Pt

RAIZ = Path(__file__).resolve().parent.parent.parent
SKILL = RAIZ / ".claude" / "skills" / "abnt-uvv-direito"
SCRIPT = SKILL / "scripts" / "verificar_docx.py"
GERADOR = RAIZ / "tests" / "checar" / "gerar_fixtures.py"


def carregar(nome, caminho):
    spec = importlib.util.spec_from_file_location(nome, caminho)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.vd = carregar("verificar_docx", SCRIPT)  # falha se o script nao existe
        cls.gf = carregar("gerar_fixtures", GERADOR)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.c01 = cls.dir / "caso-01.docx"
        cls.c03 = cls.dir / "caso-03.docx"
        cls.gf.gerar_caso01(cls.c01)
        cls.gf.gerar_caso03(cls.c03)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def itens(self, caminho):
        rel = self.vd.verificar(str(caminho))
        return {i["chave"]: i for i in rel["itens"]}

    def cli(self, caminho):
        r = subprocess.run([sys.executable, str(SCRIPT), str(caminho)], capture_output=True)
        return r.returncode, r.stdout.decode("utf-8"), r.stderr.decode("utf-8")


class TestOrigem(Base):
    def test_origem(self):
        self.assertEqual(self.vd.origem("GRA-09"), "Manual 3.2, p. 32")
        self.assertEqual(self.vd.origem("GRA-03"), "Manual 3.1, p. 32")

    def test_origem_ausente(self):
        with self.assertRaises(self.vd.RegraAusente):
            self.vd.origem("GRA-99")

    def test_trechos_dos_limiares(self):
        self.assertTrue(self.vd.LIMIARES)
        for rid, lim in self.vd.LIMIARES.items():
            bloco = self.vd.bloco_regra(rid)
            for trecho in lim["texto"]:
                self.assertIn(trecho, bloco, f"{rid}: {trecho!r}")


class TestArquivo(Base):
    def test_recusas(self):
        docm = self.dir / "a.docm"
        docm.write_bytes(b"PK\x03\x04")
        txt = self.dir / "a.txt"
        txt.write_text("x", encoding="utf-8")
        for p in (docm, txt, self.dir / "nao-existe.docx"):
            codigo, out, err = self.cli(p)
            self.assertEqual(codigo, 2, str(p))
            self.assertNotIn("Traceback", err)
            self.assertIn("Não foi possível ler o arquivo", err)

    def test_docx_falso(self):
        f = self.dir / "falso.docx"
        f.write_text("nao e zip", encoding="utf-8")
        codigo, _, err = self.cli(f)
        self.assertEqual(codigo, 2)
        self.assertNotIn("Traceback", err)


class TestCorpo(Base):
    def test_caso01(self):
        it = self.itens(self.c01)
        self.assertEqual(it["margens"]["resultado"], "✘")
        self.assertIn("2,5 cm", it["margens"]["detalhe"])
        self.assertIn("GRA-09", [r[0] for r in it["margens"]["regras"]])
        self.assertEqual(it["fonte_corpo"]["resultado"], "✘")
        self.assertIn("Times New Roman", it["fonte_corpo"]["detalhe"])
        self.assertIn("GRA-03", [r[0] for r in it["fonte_corpo"]["regras"]])
        for k in ("tamanho_corpo", "entrelinhas_corpo", "alinhamento_corpo"):
            self.assertEqual(it[k]["resultado"], "✔", k)

    def test_caso03(self):
        it = self.itens(self.c03)
        for k in ("fonte_corpo", "tamanho_corpo", "entrelinhas_corpo",
                  "alinhamento_corpo", "margens"):
            self.assertEqual(it[k]["resultado"], "✔", k)
        self.assertFalse([i for i in it.values() if i["resultado"] == "✘"])

    def test_fonte_mista(self):
        p = self.dir / "mista.docx"
        self.gf.gerar_caso03(p)
        d = Document(str(p))
        alvo = [x for x in d.paragraphs if x.text.startswith("Este paragrafo")][0]
        alvo.runs[0].font.name = "Times New Roman"
        d.save(str(p))
        it = self.itens(p)
        self.assertEqual(it["fonte_corpo"]["resultado"], "✘")
        self.assertIn("1 de", it["fonte_corpo"]["detalhe"])

    def test_mirror_margins(self):
        p = self.dir / "espelho.docx"
        self.gf.gerar_caso03(p)
        d = Document(str(p))
        d.settings.element.append(OxmlElement("w:mirrorMargins"))
        d.save(str(p))
        it = self.itens(p)
        self.assertEqual(it["margens"]["resultado"], "verificar manualmente")
        self.assertIn("GRA-10", [r[0] for r in it["margens"]["regras"]])


class TestCLI(Base):
    def test_caso01(self):
        codigo, out, _ = self.cli(self.c01)
        self.assertEqual(codigo, 1)
        self.assertIn("| # | Item | Resultado | Regra | Detalhe |", out)
        self.assertIn("✘", out)
        self.assertIn("GRA-09 (Manual 3.2, p. 32)", out)
        self.assertIn("GRA-03 (Manual 3.1, p. 32)", out)
        self.assertIn("## Verificar manualmente", out)
        self.assertRegex(out, r"Resumo: \d+ ✔, \d+ ✘, \d+ verificar manualmente")

    def test_caso03(self):
        codigo, out, _ = self.cli(self.c03)
        self.assertEqual(codigo, 0)
        # a linha Resumo traz "0 ✘" por formato (P-18); so as demais nao podem ter ✘
        restante = [x for x in out.splitlines() if not x.startswith("Resumo:")]
        self.assertNotIn("✘", " ".join(restante))
        self.assertRegex(out, r"Resumo: \d+ ✔, 0 ✘, \d+ verificar manualmente")

    def test_nao_escreve(self):
        antes = self.c03.read_bytes()
        self.cli(self.c03)
        self.assertEqual(antes, self.c03.read_bytes())
        self.assertEqual(sorted(x.name for x in self.dir.glob("~*")), [])


def _campo_simples(par, fonte="Arial", meio_pt=20):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), fonte)
    rf.set(qn("w:hAnsi"), fonte)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(meio_pt))
    rpr.append(rf)
    rpr.append(sz)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(rpr)
    r.append(t)
    fld.append(r)
    par._p.append(fld)


def _run_campo(par, tipo=None, instr=None, texto=None, fonte="Arial", meio_pt=20):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), fonte)
    rf.set(qn("w:hAnsi"), fonte)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(meio_pt))
    rpr.append(rf)
    rpr.append(sz)
    r.append(rpr)
    if tipo:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tipo)
        r.append(fc)
    if instr:
        it = OxmlElement("w:instrText")
        it.text = instr
        r.append(it)
    if texto:
        t = OxmlElement("w:t")
        t.text = texto
        r.append(t)
    par._p.append(r)


LONGO = "Texto de teste ficticio para o corpo do documento. " * 3


class TestDoc02(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.c02 = cls.dir / "caso-02.docx"
        cls.gf.gerar_caso02(cls.c02)

    def doc_simples(self, campo=True):
        """Capa, '1 INTRODUCAO' e corpo; cabecalho com campo PAGE (Arial 10, direita)."""
        d = Document()
        d.add_paragraph("CAPA DE TESTE")
        t = d.add_paragraph("1 INTRODUÇÃO")
        t.paragraph_format.page_break_before = True
        d.add_paragraph(LONGO)
        if campo:
            hp = d.sections[0].header.paragraphs[0]
            hp.alignment = 2
            _campo_simples(hp)
        return d

    def salvar(self, d, nome):
        p = self.dir / nome
        d.save(str(p))
        return p

    def test_caso02_citacao(self):
        it = self.itens(self.c02)
        for k in ("citacao_recuo", "citacao_fonte", "citacao_entrelinhas", "citacao_aspas"):
            self.assertEqual(it[k]["resultado"], "✘", k)
            self.assertIn("CIT-06", [r[0] for r in it[k]["regras"]], k)
        self.assertRegex(it["citacao_recuo"]["detalhe"], r"2,0 cm|2 cm")

    def test_caso02_titulos(self):
        it = self.itens(self.c02)
        self.assertEqual(it["titulo_pagina_nova"]["resultado"], "✘")
        self.assertIn("2 SEGUNDA SEÇÃO DE TESTE", it["titulo_pagina_nova"]["detalhe"])
        self.assertEqual(it["titulo_grafia_niveis"]["resultado"], "✘")
        self.assertIn("1.1", it["titulo_grafia_niveis"]["detalhe"])
        self.assertEqual(it["titulo_numeracao"]["resultado"], "✘")
        self.assertIn("1.1.2", it["titulo_numeracao"]["detalhe"])
        self.assertEqual(it["titulo_primaria_formato"]["resultado"], "✔")

    def test_caso02_numeracao(self):
        it = self.itens(self.c02)
        self.assertEqual(it["numeracao_posicao"]["resultado"], "✘")
        self.assertIn("rodapé", it["numeracao_posicao"]["detalhe"])
        self.assertEqual(it["numeracao_alinhamento"]["resultado"], "✘")
        self.assertEqual(it["numeracao_fonte"]["resultado"], "✘")

    def test_caso02_nao_contamina_doc01(self):
        it = self.itens(self.c02)
        for k in ("fonte_corpo", "tamanho_corpo", "entrelinhas_corpo",
                  "alinhamento_corpo", "margens"):
            self.assertEqual(it[k]["resultado"], "✔", k)

    def test_caso02_citacao_sem_recuo(self):
        it = self.itens(self.c02)
        item = it["citacao_sem_recuo"]
        self.assertEqual(item["resultado"], "verificar manualmente")
        d = Document(str(self.c02))
        alvo = next(i for i, p in enumerate(d.paragraphs) if "Este trecho entre aspas" in p.text)
        self.assertIn(f"§{alvo}", item["detalhe"])
        self.assertIn("estimativa", item["detalhe"])

    def test_caso03_itens_novos(self):
        it = self.itens(self.c03)
        manuais = {"citacao_extensao", "citacao_sem_recuo", "numeracao_inicio"}
        for k, v in it.items():
            if k in manuais:
                self.assertEqual(v["resultado"], "verificar manualmente", k)
            else:
                self.assertEqual(v["resultado"], "✔", k)
        for k in ("citacao_recuo", "citacao_fonte", "citacao_entrelinhas", "citacao_aspas",
                  "numeracao_posicao", "numeracao_alinhamento", "numeracao_fonte",
                  "numeracao_formato", "titulo_pagina_nova", "titulo_primaria_formato",
                  "titulo_grafia_niveis", "titulo_numeracao"):
            self.assertIn(k, it)

    def test_caso01_dois_x(self):
        it = self.itens(self.c01)
        self.assertEqual(sorted(k for k, v in it.items() if v["resultado"] == "✘"),
                         ["fonte_corpo", "margens"])

    def test_manuais_em_todos(self):
        for c in (self.c01, self.c02, self.c03):
            it = self.itens(c)
            for k in ("numeracao_inicio", "citacao_extensao"):
                self.assertEqual(it[k]["resultado"], "verificar manualmente", k)

    def test_heading_sem_numero(self):
        d = Document()
        d.add_paragraph("INTRODUÇÃO", style="Heading 1")
        d.add_paragraph(LONGO)
        it = self.itens(self.salvar(d, "heading.docx"))
        self.assertEqual(it["titulo_numeracao"]["resultado"], "verificar manualmente")

    def _dois_titulos(self, tipo_secao, nome):
        d = self.doc_simples()
        d.add_section(tipo_secao)
        d.add_paragraph("2 SEGUNDA SEÇÃO DE TESTE")
        d.add_paragraph(LONGO)
        return self.itens(self.salvar(d, nome))

    def test_secao_nextpage(self):
        it = self._dois_titulos(WD_SECTION.NEW_PAGE, "np.docx")
        self.assertEqual(it["titulo_pagina_nova"]["resultado"], "✔")

    def test_secao_continua(self):
        it = self._dois_titulos(WD_SECTION.CONTINUOUS, "ct.docx")
        self.assertEqual(it["titulo_pagina_nova"]["resultado"], "✘")
        self.assertIn("2 ", it["titulo_pagina_nova"]["detalhe"])

    def test_secao_tipos_trocados(self):
        d = self.doc_simples()
        d.add_section(WD_SECTION.NEW_PAGE)
        d.sections[-1].start_type = WD_SECTION.CONTINUOUS
        d.add_paragraph("2 SEGUNDA SEÇÃO DE TESTE")
        d.add_paragraph(LONGO)
        it = self.itens(self.salvar(d, "tr.docx"))
        self.assertEqual(it["titulo_pagina_nova"]["resultado"], "✘")

    def test_quebra_de_pagina_em_run(self):
        d = self.doc_simples()
        p = d.add_paragraph("fim da primeira")
        p.add_run().add_break(WD_BREAK.PAGE)
        d.add_paragraph("2 SEGUNDA SEÇÃO DE TESTE")
        d.add_paragraph(LONGO)
        it = self.itens(self.salvar(d, "br.docx"))
        self.assertEqual(it["titulo_pagina_nova"]["resultado"], "✔")

    def test_formato_lower_roman(self):
        d = self.doc_simples()
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:fmt"), "lowerRoman")
        d.sections[0]._sectPr.append(pg)
        it = self.itens(self.salvar(d, "rom.docx"))
        self.assertEqual(it["numeracao_formato"]["resultado"], "✘")
        self.assertIn("lowerRoman", it["numeracao_formato"]["detalhe"])

    def test_sem_campo(self):
        d = self.doc_simples(campo=False)
        it = self.itens(self.salvar(d, "sc.docx"))
        self.assertEqual(it["numeracao_posicao"]["resultado"], "✘")
        self.assertIn("ausente", it["numeracao_posicao"]["detalhe"])
        for k in ("numeracao_alinhamento", "numeracao_fonte", "numeracao_formato"):
            self.assertNotIn(k, it)

    def test_campo_complexo(self):
        d = self.doc_simples(campo=False)
        hp = d.sections[0].header.paragraphs[0]
        hp.alignment = 2
        _run_campo(hp, tipo="begin")
        _run_campo(hp, instr=" PAGE ")
        _run_campo(hp, tipo="separate")
        _run_campo(hp, texto="1")
        _run_campo(hp, tipo="end")
        it = self.itens(self.salvar(d, "cx.docx"))
        self.assertEqual(it["numeracao_posicao"]["resultado"], "✔")
        self.assertEqual(it["numeracao_alinhamento"]["resultado"], "✔")
        self.assertEqual(it["numeracao_fonte"]["resultado"], "✔")

    def test_tab_antes_do_campo(self):
        d = self.doc_simples(campo=False)
        hp = d.sections[0].header.paragraphs[0]
        hp.add_run().add_tab()
        _campo_simples(hp)
        it = self.itens(self.salvar(d, "tab.docx"))
        self.assertEqual(it["numeracao_alinhamento"]["resultado"], "verificar manualmente")

    def test_sem_citacao_candidata(self):
        d = self.doc_simples()
        it = self.itens(self.salvar(d, "sci.docx"))
        self.assertEqual(it["citacao_longa"]["resultado"], "verificar manualmente")
        self.assertNotIn("citacao_recuo", it)

    def test_cli_caso02(self):
        codigo, out, _ = self.cli(self.c02)
        self.assertEqual(codigo, 1)
        for trecho in ("CIT-06 (Manual 4.1.3, p. 37)", "GRA-16 (Manual 3.4, p. 34)",
                       "GRA-12 (Manual 3.3, p. 33)", "GRA-13 (Manual 3.1, p. 32)"):
            self.assertIn(trecho, out)
        self.assertEqual(self.cli(self.c01)[1].count("| ✘ |"), 2)
        codigo, out, _ = self.cli(self.c03)
        self.assertEqual(codigo, 0)
        manual = out.split("## Verificar manualmente")[1]
        self.assertIn("GRA-11", manual)
        self.assertIn("CIT-07", manual)


if __name__ == "__main__":
    unittest.main()

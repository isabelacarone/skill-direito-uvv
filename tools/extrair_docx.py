#!/usr/bin/env python
"""Dump por paragrafo de um .docx: 'indice | texto'. Tabelas ao final."""
import argparse
import sys

from docx import Document


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--de", type=int, default=0)
    ap.add_argument("--ate", type=int, default=None)
    a = ap.parse_args(argv)
    doc = Document(a.docx)
    for i, p in enumerate(doc.paragraphs):
        if i < a.de or (a.ate is not None and i > a.ate):
            continue
        if p.text.strip():
            print(f"{i} | {p.text}")
    for n, t in enumerate(doc.tables):
        print(f"--- tabela {n}")
        for row in t.rows:
            print(" | ".join(c.text.strip() for c in row.cells))
    return 0


if __name__ == "__main__":
    sys.exit(main())
